from groq import Groq
import streamlit as st
import json
from pptx import Presentation

st.title("ChatBot :speech_balloon:")
st.sidebar.header("ChatBot")

client = Groq(api_key=st.secrets["GROQ_API_KEY"])

# Funktion zum Extrahieren von Text aus PowerPoint-Dateien
def extract_text_from_pptx(file):
    prs = Presentation(file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

# Datei-Upload in der Sidebar
uploaded_file = st.sidebar.file_uploader("PowerPoint hochladen", type=["pptx"])

# Einmalige Extraktion & System-Prompt setzen
if uploaded_file and "ppt_knowledge" not in st.session_state:
    content = extract_text_from_pptx(uploaded_file)
    system_prompt = f"Du bist ein hilfreicher Assistent. Nutze das folgende Wissen aus einer Präsentation als Kontext:\n\n{content}"
    st.session_state.ppt_knowledge = system_prompt

# Modell setzen, falls noch nicht gesetzt
if "groq_model" not in st.session_state:
    st.session_state["groq_model"] = "llama3-8b-8192"

# Nachrichten-Session initialisieren
if "messages" not in st.session_state:
    st.session_state.messages = []
    # System-Prompt nur beim allerersten Start einfügen
    if "ppt_knowledge" in st.session_state:
        st.session_state.messages.append({
            "role": "system",
            "content": st.session_state.ppt_knowledge
        })

# Bisherige Nachrichten anzeigen
for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

# Nutzer-Eingabe
if prompt := st.chat_input("How can I assist you?"):
    st.session_state.messages.append({"role": "user", "content": prompt})
    with st.chat_message("user"):
        st.markdown(prompt)

    with st.chat_message("assistant"):
        chat_completion = client.chat.completions.create(
            model=st.session_state["groq_model"],
            messages=[
                {"role": m["role"], "content": m["content"]}
                for m in st.session_state.messages
            ],
            stream=True,
        )
        response = ""

        # Streaming der Antwort
        for mess in chat_completion:
            content_part = mess.choices[0].delta.content if mess.choices[0].delta.content else ""
            response += content_part

        st.markdown(response)

    # Antwort des Chatbots speichern
    st.session_state.messages.append({"role": "assistant", "content": response})
